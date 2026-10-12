"""Capabilities are scoped to one root; no unrestricted shell is exposed."""
import os
import re
import subprocess
from pathlib import Path

from app.core.schemas import Tool


class Workspace:
    def __init__(self, root, max_bytes=5_000_000, max_chars=50_000):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.max_bytes = max_bytes
        self.max_chars = max_chars

    def path(self, path):
        # Reject both Windows and POSIX syntax regardless of the host OS.
        normalized = str(path).replace('\\', '/')
        parts = normalized.split('/')
        sensitive = {'.git', '.ssh', '.aws', '.codex', '.agents', '.gnupg', 'credentials', 'id_rsa', 'id_ed25519'}
        reserved = re.compile(r'^(con|prn|aux|nul|com[0-9]|lpt[0-9])(?:\.|$)', re.I)
        if not normalized or normalized.startswith('/') or ':' in normalized or any(
            p in ('..', '') or p.endswith(('.', ' ')) or any(ord(c) < 32 or c in '<>"|?*' for c in p) or p.lower() in sensitive or p.lower().startswith('.env') or reserved.match(p)
            for p in parts
        ):
            raise PermissionError('path is outside the approved workspace or is sensitive')
        target = (self.root / normalized).resolve()
        if not target.is_relative_to(self.root) or target == self.root:
            raise PermissionError('path escapes the approved workspace')
        # Refuse symlinks even within root, and junctions where available.
        current = self.root
        for p in parts:
            current = current / p
            if current.is_symlink() or (hasattr(current, 'is_junction') and current.is_junction()):
                raise PermissionError('symlinks and junctions are not approved workspace paths')
        return target

    def read(self, path):
        target = self.path(path)
        if target.stat().st_size > self.max_bytes:
            raise ValueError('document is too large')
        ext = target.suffix.lower()
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(target)
            if reader.is_encrypted:
                raise ValueError('encrypted PDFs must be decrypted outside the agent workspace')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
            if not text.strip():
                raise ValueError('PDF has no extractable text; OCR is required')
        elif ext == '.docx':
            from docx import Document
            doc = Document(target)
            text = '\n'.join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for row in t.rows for cell in row.cells])
        else:
            text = target.read_text(encoding='utf-8')
        if not text.strip():
            raise ValueError('document is empty')
        if len(text) > self.max_chars:
            raise ValueError('extracted document is too large; split it into smaller files')
        return text

    def write(self, path, content):
        target = self.path(path)
        if len(content.encode('utf-8')) > self.max_bytes:
            raise ValueError('output is too large')
        target.parent.mkdir(parents=True, exist_ok=True)
        # Recheck after making parents; external processes must not mutate paths concurrently.
        target = self.path(path)
        # Preserve the requested UTF-8 bytes on Windows too (no CRLF conversion).
        target.write_bytes(content.encode('utf-8'))
        return f'Wrote {len(content)} characters to {path}'

    def _git(self, arguments):
        # Prevent inherited Git injection and optional index writes on read-only operations.
        env = {k:v for k,v in os.environ.items() if not k.startswith('GIT_')}
        env['GIT_OPTIONAL_LOCKS'] = '0'
        root = subprocess.run(['git','rev-parse','--show-toplevel'],cwd=self.root,env=env,capture_output=True,text=True,timeout=15,check=True).stdout.strip()
        if Path(root).resolve() != self.root:
            raise PermissionError('Git repository root must equal the approved workspace')
        result = subprocess.run(['git','--literal-pathspecs','--no-pager','-c','core.fsmonitor=false',*arguments],cwd=self.root,env=env,capture_output=True,text=True,timeout=15,check=True)
        return result.stdout[:self.max_chars]

    def git_status(self):
        return self._git(['status','--short','--untracked-files=normal'])

    def git_diff(self):
        changed = self._git(['diff', '--name-only', '-z', '--no-ext-diff', '--no-textconv'])
        paths = []
        for path in changed.split('\0'):
            if not path:
                continue
            try:
                self.path(path)
            except PermissionError:
                continue
            paths.append(path)
        if not paths:
            return ''
        return self._git(['diff', '--no-ext-diff', '--no-textconv', '--', *paths])


def builtin_tools():
    def parameters(properties):
        return {'type':'object','properties':{k:{'type':'string','description':v} for k,v in properties.items()},'required':list(properties),'additionalProperties':False}
    return [
        Tool(id='document_read',description='Read TXT, Markdown, PDF or DOCX inside the approved workspace',parameters=parameters({'path':'Workspace-relative document path'})),
        Tool(id='filesystem_read',description='Read a UTF-8 source or text file inside the approved workspace',parameters=parameters({'path':'Workspace-relative file path'})),
        Tool(id='filesystem_write',description='Prepare a UTF-8 file inside the approved workspace',permission_level=2,parameters=parameters({'path':'Workspace-relative file path','content':'Full file contents'})),
        Tool(id='git_status',description='Read-only Git status for the approved workspace'),
        Tool(id='git_diff',description='Read-only Git diff for the approved workspace'),
    ]


class ToolExecutor:
    def __init__(self, workspace):
        self.workspace = workspace
        self.handlers = {
            'document_read':workspace.read, 'filesystem_read':workspace.read,
            'filesystem_write':workspace.write, 'git_status':workspace.git_status,
            'git_diff':workspace.git_diff,
        }

    def register(self, tool_id, handler):
        """Trusted application plugins can register a handler without changing agents."""
        self.handlers[tool_id] = handler

    def execute(self, agent, tool, arguments):
        if tool.id not in agent.tools or not tool.enabled or agent.permission_level < tool.permission_level:
            raise PermissionError(f'tool {tool.id} is not permitted for agent {agent.id}')
        if tool.id not in self.handlers:
            raise ValueError(f'no trusted implementation registered for tool {tool.id}')
        schema = tool.parameters
        properties = schema.get('properties', {})
        if not isinstance(arguments, dict) or set(arguments) - set(properties) or set(schema.get('required', [])) - set(arguments):
            raise ValueError('invalid tool arguments')
        if any(not isinstance(v,str) for v in arguments.values()):
            raise ValueError('tool arguments must be strings')
        return self.handlers[tool.id](**arguments)
